from __future__ import annotations

from conftest import add_student
from openpyxl import load_workbook

from moove_recovery.infrastructure.reports import (
    export_monthly_report_pdf,
    export_monthly_report_xlsx,
)


def _report_with_voided_payment(environment) -> dict[str, object]:
    service = environment["service"]
    owner = environment["owner"]
    student_id = add_student(environment, name="=1+1", dni="0012345678")
    service.generate_missing_dues()
    student = next(item for item in service.list_students(owner) if item["id"] == student_id)
    fee = student["fees"][0]
    payment_id = service.record_payment(
        owner,
        student_id=student_id,
        due_ids=[fee["id"]],
        total_cents=fee["amount_cents"],
        method="transferencia",
        reference="TRANSFER-001",
    )
    service.void_payment(owner, payment_id, "Cobro duplicado")
    return service.get_monthly_report(owner, 2026, 10)


def test_monthly_report_includes_balances_methods_and_voids(environment) -> None:
    report = _report_with_voided_payment(environment)
    summary = report["summary"]
    fee = report["dues"][0]
    payment = report["payments"][0]

    assert summary["received_cents"] == 0
    assert fee["student_name"] == "=1+1 Pérez"
    assert fee["outstanding_cents"] == fee["amount_cents"]
    assert fee["phone"] == "011 4444-0000"
    assert payment["method"] == "transferencia"
    assert payment["status"] == "voided"
    assert payment["void_reason"] == "Cobro duplicado"
    assert payment["periods"] == "2026-10"


def test_xlsx_export_keeps_ids_as_text_and_amounts_numeric(environment, tmp_path) -> None:
    report = _report_with_voided_payment(environment)
    path = tmp_path / "informe.xlsx"
    export_monthly_report_xlsx(report, path)

    workbook = load_workbook(path, data_only=False)
    summary = workbook["Resumen"]
    dues = workbook["Cuotas"]
    payments = workbook["Cobros"]
    assert summary["B4"].value == 0
    assert summary["B7"].value == 30000
    assert dues["A2"].value == "=1+1 Pérez"
    assert dues["A2"].data_type == "s"
    assert dues["B2"].value == "0012345678"
    assert dues["B2"].data_type == "s"
    assert dues["C2"].value.startswith("011")
    assert dues["C2"].data_type == "s"
    assert dues["G2"].value == 30000
    assert isinstance(dues["G2"].value, (int, float))
    assert payments["G2"].value == "Anulado"
    assert payments["H2"].value == "Cobro duplicado"


def test_pdf_and_empty_period_exports_are_valid(environment, tmp_path) -> None:
    service = environment["service"]
    populated_report = _report_with_voided_payment(environment)
    empty_report = service.get_monthly_report(environment["owner"], 2026, 11)
    assert empty_report["dues"] == []
    assert empty_report["payments"] == []

    populated_pdf = tmp_path / "informe-con-datos.pdf"
    empty_pdf = tmp_path / "informe-vacio.pdf"
    xlsx_path = tmp_path / "informe-vacio.xlsx"
    export_monthly_report_pdf(populated_report, populated_pdf)
    export_monthly_report_pdf(empty_report, empty_pdf)
    export_monthly_report_xlsx(empty_report, xlsx_path)

    assert populated_pdf.read_bytes().startswith(b"%PDF-")
    assert populated_pdf.stat().st_size > 1000
    assert empty_pdf.read_bytes().startswith(b"%PDF-")
    assert empty_pdf.stat().st_size > 1000
    workbook = load_workbook(xlsx_path, data_only=True)
    assert workbook["Resumen"]["B4"].value == 0
    assert workbook["Cuotas"].max_row == 1
    assert workbook["Cobros"].max_row == 1
