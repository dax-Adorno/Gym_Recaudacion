import sqlite3
from datetime import date

import pytest
from conftest import add_student
from openpyxl import load_workbook
from PySide6.QtWidgets import QMessageBox

from moove_recovery.domain.errors import PermissionDenied, ValidationError
from moove_recovery.infrastructure.backups import BackupManager
from moove_recovery.infrastructure.database import MIGRATION_1, Database
from moove_recovery.infrastructure.reports import (
    export_monthly_report_pdf,
    export_monthly_report_xlsx,
)
from moove_recovery.ui.expenses_page import ExpenseDialog, ExpensesPage
from moove_recovery.ui.main_window import MainWindow


def add_expense(environment, **changes):
    data = dict(
        concept="Electricidad",
        category="Servicios",
        scheduled_on=date(2026, 10, 2),
        amount_cents=123456,
        paid_on=date(2026, 10, 2),
        method="efectivo",
    )
    data.update(changes)
    return environment["service"].save_expense(environment["owner"], **data)


def test_paid_pending_and_month_of_payment(environment):
    service, owner = environment["service"], environment["owner"]
    add_expense(environment)
    pending = add_expense(environment, paid_on=None, method=None)
    summary = service.get_dashboard_summary(owner, 2026, 10)
    assert summary["expense_cents"] == 123456
    assert summary["pending_expense_cents"] == 123456
    assert summary["net_cents"] == -123456
    environment["clock"].current = date(2026, 11, 5)
    service.pay_expense(owner, pending, paid_on=date(2026, 11, 5), method="transferencia")
    assert service.get_dashboard_summary(owner, 2026, 10)["expense_cents"] == 123456
    assert service.get_dashboard_summary(owner, 2026, 11)["expense_cents"] == 123456
    assert service.get_dashboard_summary(owner, 2026, 11)["pending_expense_cents"] == 0
    with pytest.raises(ValidationError):
        service.pay_expense(owner, pending, paid_on=date(2026, 11, 5), method="efectivo")


@pytest.mark.parametrize(
    "changes",
    [
        {"amount_cents": 0},
        {"amount_cents": -10},
        {"amount_cents": 1.5},
        {"amount_cents": True},
        {"concept": " "},
        {"category": ""},
        {"method": "inventado"},
        {"paid_on": date(2026, 11, 1)},
        {"paid_on": None},
    ],
)
def test_invalid_expense_leaves_no_changes(environment, changes):
    with pytest.raises(ValidationError):
        add_expense(environment, **changes)
    assert environment["db"].fetch_one("SELECT COUNT(*) FROM expenses")[0] == 0


def test_edit_void_and_audit(environment):
    service, owner = environment["service"], environment["owner"]
    expense = add_expense(environment)
    add_expense(environment, expense_id=expense, amount_cents=222200)
    assert service.get_dashboard_summary(owner, 2026, 10)["expense_cents"] == 222200
    service.void_expense(owner, expense, "Error de carga")
    assert service.get_dashboard_summary(owner, 2026, 10)["expense_cents"] == 0
    assert service.list_expenses(owner, status="voided")[0]["void_reason"] == "Error de carga"
    with pytest.raises(ValidationError):
        service.void_expense(owner, expense, "Otra vez")
    with pytest.raises(ValidationError):
        add_expense(environment, expense_id=expense)
    actions = environment["db"].fetch_all(
        "SELECT action FROM audit_log WHERE entity_type='expenses'"
    )
    assert [row[0] for row in actions] == ["gasto_alta", "gasto_edicion", "gasto_anulacion"]


def test_employee_cannot_access_expenses(environment, qtbot):
    service, owner = environment["service"], environment["owner"]
    service.create_employee(
        owner, full_name="Empleado Ficticio", username="empleado", password="ClaveTemporal-123"
    )
    employee = service.login("empleado", "ClaveTemporal-123")
    expense = add_expense(environment)
    for call in [
        lambda: service.list_expenses(employee, status="paid"),
        lambda: service.save_expense(
            employee, concept="x", category="x", scheduled_on=date(2026, 10, 2), amount_cents=100
        ),
        lambda: service.pay_expense(
            employee, expense, paid_on=date(2026, 10, 2), method="efectivo"
        ),
        lambda: service.void_expense(employee, expense, "x"),
        lambda: service.get_monthly_report(employee, 2026, 10),
    ]:
        with pytest.raises(PermissionDenied):
            call()
    window = MainWindow(service, employee)
    qtbot.addWidget(window)
    assert "Gastos" not in window.page_widgets
    with pytest.raises(PermissionDenied):
        ExpensesPage(service, employee)


def test_expense_exports_totals_and_safe_text(environment, tmp_path):
    add_expense(environment, concept="=1+1")
    add_expense(environment, paid_on=None, method=None)
    report = environment["service"].get_monthly_report(environment["owner"], 2026, 10)
    path = tmp_path / "economia.xlsx"
    export_monthly_report_xlsx(report, path)
    book = load_workbook(path)
    assert book["Resumen"]["B9"].value == 1234.56
    assert book["Resumen"]["B10"].value == -1234.56
    assert book["Gastos pagados"]["B2"].data_type == "s"
    assert book["Gastos pendientes"].max_row == 2
    pdf = tmp_path / "economia.pdf"
    export_monthly_report_pdf(report, pdf)
    assert pdf.read_bytes().startswith(b"%PDF-")


def test_schema_one_migrates_and_backups(tmp_path):
    path = tmp_path / "old.sqlite3"
    connection = sqlite3.connect(path)
    connection.executescript(MIGRATION_1 + "PRAGMA user_version=1;")
    connection.execute("INSERT INTO settings VALUES ('preserved','yes')")
    connection.commit()
    connection.close()
    db = Database(path)
    db.migrate()
    assert db.fetch_one("PRAGMA user_version")[0] == 2
    assert db.fetch_one("SELECT value FROM settings WHERE key='preserved'")[0] == "yes"
    backups = list((tmp_path / "backups").glob("pre-migration*.sqlite3"))
    assert len(backups) == 1
    db.migrate()
    assert db.fetch_one("SELECT COUNT(*) FROM expenses")[0] == 0


def test_expense_dialog_and_tables(environment, qtbot):
    service, owner = environment["service"], environment["owner"]
    dialog = ExpenseDialog(service, owner, pending=True)
    qtbot.addWidget(dialog)
    dialog.concept.setText("Limpieza")
    dialog.amount.setText("1234,56")
    dialog.save()
    assert dialog.result() == dialog.DialogCode.Accepted
    page = ExpensesPage(service, owner)
    qtbot.addWidget(page)
    page.show()
    page.tabs.setCurrentIndex(1)
    page.tables[1].selectRow(0)
    assert page.pay.isEnabled()
    expense = page.selected()
    payment = ExpenseDialog(service, owner, expense, pay_only=True)
    qtbot.addWidget(payment)
    payment.save()
    page.refresh()
    assert page.tables[0].rowCount() == 1
    assert page.tables[1].rowCount() == 0


def test_net_uses_cash_received_and_excludes_voided_payments(environment):
    service, owner = environment["service"], environment["owner"]
    student_id = add_student(environment)
    service.generate_missing_dues()
    fee = service.list_students(owner)[0]["fees"][0]
    payment = service.record_payment(
        owner,
        student_id=student_id,
        due_ids=[fee["id"]],
        total_cents=fee["amount_cents"],
        method="efectivo",
    )
    add_expense(environment)
    summary = service.get_dashboard_summary(owner, 2026, 10)
    assert summary["net_cents"] == fee["amount_cents"] - 123456
    service.void_payment(owner, payment, "Error de carga")
    assert service.get_dashboard_summary(owner, 2026, 10)["net_cents"] == -123456


def test_expense_payment_rolls_back_when_audit_fails(environment):
    service, owner = environment["service"], environment["owner"]
    expense = add_expense(environment, paid_on=None, method=None)
    with environment["db"].transaction() as connection:
        connection.execute(
            "CREATE TRIGGER reject_expense_audit BEFORE INSERT ON audit_log "
            "WHEN NEW.action='gasto_pago' BEGIN SELECT RAISE(ABORT,'test'); END"
        )
    with pytest.raises(sqlite3.IntegrityError):
        service.pay_expense(owner, expense, paid_on=date(2026, 10, 3), method="efectivo")
    assert service.list_expenses(owner, status="pending")[0]["id"] == expense


def test_backup_restores_expenses_and_accepts_schema_one(environment, tmp_path):
    service, owner = environment["service"], environment["owner"]
    expense = add_expense(environment)
    manager = BackupManager(environment["db"])
    backup = manager.create_backup(tmp_path / "with-expenses.sqlite3")
    service.void_expense(owner, expense, "Error")
    manager.restore_backup(backup)
    assert service.list_expenses(owner, status="paid")[0]["id"] == expense
    old = tmp_path / "schema-one.sqlite3"
    environment["db"].backup_to(old)
    connection = sqlite3.connect(old)
    connection.execute("DROP TABLE expenses")
    connection.execute("PRAGMA user_version=1")
    connection.commit()
    connection.close()
    manager.restore_backup(old)
    assert environment["db"].fetch_one("PRAGMA user_version")[0] == 2
    assert service.list_expenses(owner, status="paid") == []
    assert service.is_setup_complete()


def test_expenses_calendar_uses_payment_date_and_marks_pending_and_voided(environment, qtbot):
    from moove_recovery.ui.calendar_page import CalendarPage

    service, owner = environment["service"], environment["owner"]
    paid = add_expense(environment, scheduled_on=date(2026, 9, 28))
    pending = add_expense(environment, paid_on=None, method=None, scheduled_on=date(2026, 10, 20))
    events = service.list_calendar_events(owner, date(2026, 10, 1), date(2026, 10, 31), "gastos")
    assert [(item["expense_id"], item["date"]) for item in events] == [
        (paid, date(2026, 10, 2)),
        (pending, date(2026, 10, 20)),
    ]
    assert "pagado" in events[0]["details"]
    assert "pendiente" in events[1]["details"]
    service.void_expense(owner, paid, "Error de carga")
    assert (
        "anulado"
        in service.list_calendar_events(owner, date(2026, 10, 1), date(2026, 10, 31), "gastos")[0][
            "details"
        ]
    )
    page = CalendarPage(service, owner)
    qtbot.addWidget(page)
    page.select_category("gastos")
    assert page.category_buttons["gastos"].isChecked()
    page.selected_date = date(2026, 10, 20)
    page.refresh()
    assert page.events_table.rowCount() == 1
    assert "pendiente" in page.events_table.item(0, 1).text()
    assert "gastos" in page.category_buttons


def test_expense_table_fills_space_after_skewed_column_resize(environment, qtbot):
    page = ExpensesPage(environment["service"], environment["owner"])
    qtbot.addWidget(page)
    page.resize(1200, 600)
    page.show()
    table = page.tables[0]
    table.horizontalHeader().resizeSection(1, 700)
    page.resize(950, 600)
    qtbot.waitUntil(lambda: table.horizontalHeader().length() == table.viewport().width())
    assert all(table.columnWidth(index) >= 95 for index in range(6))


def test_edit_expense_status_recalculates_and_preserves_audit(environment):
    service, owner = environment["service"], environment["owner"]
    expense = add_expense(environment)
    add_expense(environment, expense_id=expense, paid_on=None, method=None, expected_status="paid")
    row = service.list_expenses(owner, status="pending")[0]
    assert row["paid_on"] is None and row["method"] is None
    summary = service.get_dashboard_summary(owner, 2026, 10)
    assert summary["expense_cents"] == 0
    assert summary["pending_expense_cents"] == 123456
    add_expense(environment, expense_id=expense, expected_status="pending")
    assert service.get_dashboard_summary(owner, 2026, 10)["expense_cents"] == 123456
    assert environment["db"].fetch_one("SELECT COUNT(*) FROM expenses")[0] == 1
    audits = environment["db"].fetch_all(
        "SELECT details FROM audit_log WHERE action='gasto_edicion'"
    )
    assert len(audits) == 2
    assert '"paid_on": "2026-10-02"' in audits[0][0]


def test_expense_state_edit_cancel_and_confirm(environment, qtbot, monkeypatch):
    service, owner = environment["service"], environment["owner"]
    add_expense(environment)
    expense = service.list_expenses(owner, status="paid")[0]
    dialog = ExpenseDialog(service, owner, expense)
    qtbot.addWidget(dialog)
    assert dialog.status.isEnabled()
    dialog.status.setCurrentIndex(1)
    assert not dialog.paid.isEnabled() and not dialog.method.isEnabled()
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.StandardButton.No)
    dialog.save()
    assert len(service.list_expenses(owner, status="paid")) == 1
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.StandardButton.Yes)
    dialog.save()
    assert dialog.result() == dialog.DialogCode.Accepted
    assert len(service.list_expenses(owner, status="pending")) == 1


def test_stale_expense_editor_does_not_reverse_new_payment(environment):
    service, owner = environment["service"], environment["owner"]
    expense = add_expense(environment, paid_on=None, method=None)
    service.pay_expense(owner, expense, paid_on=date(2026, 10, 3), method="transferencia")
    with pytest.raises(ValidationError, match="estado del gasto"):
        add_expense(
            environment, expense_id=expense, paid_on=None, method=None, expected_status="pending"
        )
    assert len(service.list_expenses(owner, status="paid")) == 1
