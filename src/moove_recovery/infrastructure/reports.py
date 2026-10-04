from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from html import escape
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ReportData = dict[str, Any]
MONEY_FORMAT = '"$" #,##0.00'
HEADER_FILL = "173F35"


def export_monthly_report_xlsx(report: ReportData, path: Path) -> None:
    summary = report["summary"]
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Resumen"
    overview.append(["MOOVE RECOVERY · Gestión de cuotas"])
    overview.append(["Período", summary["period"]])
    overview.append(["Fecha de corte", summary["cutoff_date"]])
    overview.append(["Recibido por fecha de cobro", _excel_amount(summary["received_cents"])])
    overview.append(["Aplicado al período", _excel_amount(summary["applied_cents"])])
    overview.append(["Pendiente del período", _excel_amount(summary["pending_cents"])])
    overview.append(["Total final de cuotas", _excel_amount(summary["period_total_cents"])])
    overview.append(["Cantidad de cuotas", int(summary["due_count"])])
    for row in range(4, 8):
        overview.cell(row=row, column=2).number_format = MONEY_FORMAT
    overview.column_dimensions["A"].width = 38
    overview.column_dimensions["B"].width = 24
    _style_header(overview, 1, 2)

    dues = workbook.create_sheet("Cuotas")
    dues_headers = [
        "Alumno",
        "DNI",
        "Teléfono",
        "Estado alumno",
        "Período",
        "Vencimiento",
        "Precio base",
        "Descuento",
        "Origen descuento",
        "Importe cuota",
        "Pagado válido",
        "Saldo",
    ]
    dues.append(dues_headers)
    for fee in report["dues"]:
        dues.append(
            [
                fee["student_name"],
                str(fee["dni"]),
                str(fee["phone"]),
                "Activo" if fee["active"] else "Inactivo",
                fee["period"],
                _excel_date(fee["due_date"]),
                _excel_amount(fee["base_cents"]),
                _excel_amount(fee["discount_cents"]),
                _discount_label(fee["discount_source"]),
                _excel_amount(fee["amount_cents"]),
                _excel_amount(fee["paid_cents"]),
                _excel_amount(fee["outstanding_cents"]),
            ]
        )
        row = dues.max_row
        for column in (7, 8, 10, 11, 12):
            dues.cell(row, column).number_format = MONEY_FORMAT
        for column in (2, 3):
            dues.cell(row, column).number_format = "@"
            dues.cell(row, column).data_type = "s"
        dues.cell(row, 6).number_format = "dd/mm/yyyy"
    _style_table(dues, [30, 16, 18, 14, 12, 16, 17, 17, 20, 17, 17, 17])

    payments = workbook.create_sheet("Cobros")
    payments.append(
        [
            "ID",
            "Fecha de cobro",
            "Importe",
            "Medio",
            "Períodos cubiertos",
            "Registrado por",
            "Estado",
            "Motivo de anulación",
            "Referencia",
            "Recibido en período",
        ]
    )
    for payment in report["payments"]:
        payments.append(
            [
                int(payment["id"]),
                _excel_datetime(payment["paid_at"]),
                _excel_amount(payment["total_cents"]),
                _payment_method(payment["method"]),
                str(payment["periods"] or ""),
                payment["created_by"],
                "Anulado" if payment["status"] == "voided" else "Válido",
                str(payment["void_reason"] or ""),
                str(payment["reference"] or ""),
                "Sí" if payment["paid_in_month"] else "No",
            ]
        )
        payments.cell(payments.max_row, 3).number_format = MONEY_FORMAT
        payments.cell(payments.max_row, 2).number_format = "dd/mm/yyyy hh:mm"
    _style_table(payments, [9, 22, 16, 18, 27, 28, 12, 32, 24, 18])

    workbook.properties.title = f"MOOVE RECOVERY · Informe {summary['period']}"
    workbook.properties.subject = "Cuotas y cobranzas"
    workbook.save(path)


def export_monthly_report_pdf(report: ReportData, path: Path) -> None:
    summary = report["summary"]
    page_size = landscape(A4)
    left_margin = 12 * mm
    right_margin = 12 * mm
    top_margin = 19 * mm
    bottom_margin = 14 * mm
    doc = SimpleDocTemplate(
        str(path),
        pagesize=page_size,
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin,
        title=f"MOOVE RECOVERY · Informe {summary['period']}",
        author="MOOVE RECOVERY",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=21,
            textColor=colors.HexColor("#19384A"),
            alignment=TA_LEFT,
            spaceAfter=3 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            "ReportSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#276B4C"),
            spaceBefore=4 * mm,
            spaceAfter=2 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            "ReportSmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7,
            leading=8,
            spaceAfter=0,
        )
    )
    story: list[Any] = [
        Paragraph("MOOVE RECOVERY · Gestión de cuotas", styles["ReportTitle"]),
        Paragraph(
            f"Período {escape(str(summary['period']))} · Fecha de corte "
            f"{escape(str(summary['cutoff_date']))} · estado actual del período",
            styles["BodyText"],
        ),
        Spacer(1, 4 * mm),
    ]

    summary_table = Table(
        [
            [
                "Recibido en el mes",
                "Aplicado al período",
                "Pendiente del período",
                "Total final de cuotas",
            ],
            [
                _pesos(summary["received_cents"]),
                _pesos(summary["applied_cents"]),
                _pesos(summary["pending_cents"]),
                _pesos(summary["period_total_cents"]),
            ],
        ],
        colWidths=[(page_size[0] - left_margin - right_margin) / 4] * 4,
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F1ED")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#31433A")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D8E2DC")),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D8E2DC")),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([summary_table, Paragraph("Cuotas del período", styles["ReportSection"])])

    due_headers = [
        "Alumno",
        "DNI",
        "Estado",
        "Período",
        "Vence",
        "Base",
        "Descuento",
        "Origen",
        "Cuota final",
        "Pagado",
        "Saldo",
    ]
    due_rows = [[_paragraph(value, styles) for value in due_headers]]
    for fee in report["dues"]:
        due_rows.append(
            [
                _paragraph(fee["student_name"], styles),
                _paragraph(fee["dni"], styles),
                "Activo" if fee["active"] else "Inactivo",
                fee["period"],
                _short_date(fee["due_date"]),
                _pesos(fee["base_cents"]),
                _pesos(fee["discount_cents"]),
                _discount_label(fee["discount_source"]),
                _pesos(fee["amount_cents"]),
                _pesos(fee["paid_cents"]),
                _pesos(fee["outstanding_cents"]),
            ]
        )
    if len(due_rows) == 1:
        story.append(Paragraph("Sin cuotas registradas para este período.", styles["BodyText"]))
    else:
        due_widths = [
            86,
            48,
            42,
            43,
            46,
            56,
            58,
            58,
            60,
            54,
            54,
        ]
        story.append(_report_table(due_rows, due_widths, numeric_from=5))

    story.append(Paragraph("Cobros relacionados", styles["ReportSection"]))
    payment_headers = [
        "Fecha",
        "Importe",
        "Medio",
        "Períodos cubiertos",
        "Registrado por",
        "Estado",
        "Motivo de anulación",
        "Referencia",
        "Pago en mes",
    ]
    payment_rows = [[_paragraph(value, styles) for value in payment_headers]]
    for payment in report["payments"]:
        payment_rows.append(
            [
                _short_datetime(payment["paid_at"]),
                _pesos(payment["total_cents"]),
                _payment_method(payment["method"]),
                _paragraph(payment["periods"] or "", styles),
                _paragraph(payment["created_by"], styles),
                "Anulado" if payment["status"] == "voided" else "Válido",
                _paragraph(payment["void_reason"] or "", styles),
                _paragraph(payment["reference"] or "", styles),
                "Sí" if payment["paid_in_month"] else "No",
            ]
        )
    if len(payment_rows) == 1:
        story.append(Paragraph("Sin cobros relacionados con este período.", styles["BodyText"]))
    else:
        payment_widths = [
            69,
            56,
            57,
            88,
            91,
            43,
            105,
            84,
            65,
        ]
        story.append(_report_table(payment_rows, payment_widths, numeric_from=1))

    def draw_page(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D8E2DC"))
        canvas.line(
            left_margin, page_size[1] - 12 * mm, page_size[0] - right_margin, page_size[1] - 12 * mm
        )
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(colors.HexColor("#19384A"))
        canvas.drawString(left_margin, page_size[1] - 9 * mm, "MOOVE RECOVERY")
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#65736A"))
        canvas.drawRightString(page_size[0] - right_margin, 7 * mm, f"Página {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_page, onLaterPages=draw_page)


def _style_header(sheet: Worksheet, row: int, columns: int) -> None:
    for cell in sheet[row][:columns]:
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)


def _style_table(sheet: Worksheet, widths: list[int]) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def _excel_amount(cents: object) -> float:
    return float(Decimal(int(cents)) / Decimal(100))


def _excel_date(value: object) -> datetime:
    return datetime.strptime(str(value), "%Y-%m-%d")


def _excel_datetime(value: object) -> datetime:
    parsed = datetime.fromisoformat(str(value))
    return parsed.replace(tzinfo=None)


def _pesos(cents: object) -> str:
    formatted = f"{Decimal(int(cents)) / Decimal(100):,.2f}"
    return "$ " + formatted.replace(",", "_").replace(".", ",").replace("_", ".")


def _short_date(value: object) -> str:
    return datetime.strptime(str(value), "%Y-%m-%d").strftime("%d/%m/%Y")


def _short_datetime(value: object) -> str:
    return datetime.fromisoformat(str(value)).strftime("%d/%m/%Y %H:%M")


def _discount_label(value: object) -> str:
    labels = {"alta_50": "Alta 50 %"}
    return labels.get(str(value), str(value)) if value else "Sin descuento"


def _payment_method(value: object) -> str:
    labels = {
        "efectivo": "Efectivo",
        "transferencia": "Transferencia",
        "debito": "Débito",
        "credito": "Crédito",
        "otro": "Otro",
    }
    return labels.get(str(value), str(value))


def _paragraph(value: object, styles: Any) -> Paragraph:
    clean = str(value).encode("cp1252", "replace").decode("cp1252")
    return Paragraph(escape(clean), styles["ReportSmall"])


def _report_table(rows: list[list[Any]], widths: list[int], *, numeric_from: int) -> LongTable:
    table = LongTable(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173F35")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 6.7),
                ("LEADING", (0, 0), (-1, -1), 7.5),
                ("ALIGN", (numeric_from, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8E2DC")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9F7")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table
